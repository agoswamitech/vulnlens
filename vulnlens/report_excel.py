"""Write the analysis to a formatted Excel workbook with native (editable) charts."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from .analysis import Report
from .scoring import SLA_DAYS, W_ASSET, W_SEVERITY, W_THREAT

NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=NAVY)
SUBTITLE_FONT = Font(italic=True, size=10, color="595959")
THIN = Side(style="thin", color="D9D9D9")
PRIORITY_FILLS = {
    "P1": PatternFill("solid", fgColor="F4B6B6"),
    "P2": PatternFill("solid", fgColor="F8CBAD"),
    "P3": PatternFill("solid", fgColor="FFE699"),
    "P4": PatternFill("solid", fgColor="C6E0B4"),
}

FINDING_COLUMNS = [
    ("finding_id", "Finding"), ("priority", "Priority"), ("risk_score", "Risk Score"), ("hostname", "Hostname"),
    ("business_unit", "Business Unit"), ("owner", "Owner"), ("vuln_id", "Vuln ID"), ("title", "Title"),
    ("severity", "CVSS Severity"), ("cvss_base", "CVSS"), ("epss", "EPSS"), ("known_exploited", "Known Exploited"),
    ("internet_facing", "Internet Facing"), ("data_classification", "Data Class"), ("environment", "Environment"),
    ("first_seen", "First Seen"), ("due_date", "SLA Due"), ("age_days", "Age (days)"), ("days_past_due", "Days Past Due"),
]


def _write_table(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1, index: bool = False) -> tuple[int, int]:
    """Write df with a styled header; return (last_row, last_col)."""
    rows = list(dataframe_to_rows(df, index=index, header=True))
    if index:
        rows = [rows[0]] + rows[2:]  # dataframe_to_rows emits an extra index-name row
    for r_off, row in enumerate(rows):
        for c_off, val in enumerate(row):
            cell = ws.cell(row=start_row + r_off, column=start_col + c_off, value=val)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if r_off == 0:
                cell.fill, cell.font = HEADER_FILL, HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if index and df.index.name:
        ws.cell(row=start_row, column=start_col, value=df.index.name)
    return start_row + len(rows) - 1, start_col + len(rows[0]) - 1


def _autosize(ws, min_width: int = 8, max_width: int = 50) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font, ws["A2"].font = TITLE_FONT, SUBTITLE_FONT


def _bar(title: str, ws, min_col: int, min_row: int, max_col: int, max_row: int, stacked: bool = False, horizontal: bool = False) -> BarChart:
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title, ch.height, ch.width = title, 7.5, 15
    if stacked:
        ch.grouping, ch.overlap = "stacked", 100
    data = Reference(ws, min_col=min_col + 1, min_row=min_row, max_col=max_col, max_row=max_row)
    cats = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    return ch


def summary_sheet(wb: Workbook, r: Report) -> None:
    ws = wb.active
    ws.title = "Executive Summary"
    _title(ws, "Vulnerability Risk Report", f"As of {r.as_of:%B %d, %Y} · synthetic data for demonstration")

    ws["A4"] = "Key Metrics"
    ws["A4"].font = Font(bold=True, size=12, color=NAVY)
    kpi_df = pd.DataFrame(list(r.kpis.items()), columns=["Metric", "Value"])
    last, _ = _write_table(ws, kpi_df, start_row=5)

    row = last + 2
    sev_end, _ = _write_table(ws, r.by_severity, start_row=row)
    ws.add_chart(_bar("Open Findings by CVSS Severity", ws, 1, row, 2, sev_end), "G4")

    row = sev_end + 2
    pri = r.by_priority[["Priority", "Open", "Overdue"]]
    pri_end, _ = _write_table(ws, pri, start_row=row)
    ws.add_chart(_bar("Open vs Overdue by Risk Priority", ws, 1, row, 3, pri_end), "G20")

    row = pri_end + 2
    trend_end, _ = _write_table(ws, r.trend, start_row=row)
    line = LineChart()
    line.title, line.height, line.width = "Monthly Trend: Opened vs Closed vs Backlog", 7.5, 15
    line.add_data(Reference(ws, min_col=2, min_row=row, max_col=4, max_row=trend_end), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=row + 1, max_row=trend_end))
    ws.add_chart(line, "G36")
    _autosize(ws)
    ws.column_dimensions["A"].width = 42


def findings_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(name)
    cols = [c for c, _ in FINDING_COLUMNS]
    out = df[cols].rename(columns=dict(FINDING_COLUMNS)).sort_values(["Risk Score"], ascending=False)
    for c in ("First Seen", "SLA Due"):
        out[c] = out[c].dt.date
    last_row, last_col = _write_table(ws, out)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"
    for p, fill in PRIORITY_FILLS.items():
        ws.conditional_formatting.add(f"B2:B{last_row}", CellIsRule(operator="equal", formula=[f'"{p}"'], fill=fill))
    ws.conditional_formatting.add(
        f"C2:C{last_row}", ColorScaleRule(start_type="num", start_value=20, start_color="63BE7B", mid_type="num", mid_value=55, mid_color="FFEB84", end_type="num", end_value=90, end_color="F8696B")
    )
    _autosize(ws)


def simple_sheet(wb: Workbook, name: str, title: str, df: pd.DataFrame, note: str = "", index: bool = False, chart: dict | None = None) -> None:
    ws = wb.create_sheet(name)
    _title(ws, title, note)
    last_row, last_col = _write_table(ws, df, start_row=4, index=index)
    if chart:
        ws.add_chart(_bar(chart["title"], ws, 1, 4, chart.get("max_col", last_col), last_row, chart.get("stacked", False), chart.get("horizontal", False)), chart.get("anchor", f"{get_column_letter(last_col + 2)}4"))
    _autosize(ws)


def methodology_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Methodology")
    _title(ws, "Scoring Methodology", "How each finding's risk score and priority are calculated")
    lines = [
        ("Risk Score (0-100)", f"100 × ({W_SEVERITY:.2f} × CVSS/10 + {W_THREAT:.2f} × Threat + {W_ASSET:.2f} × Asset)"),
        ("Threat (0-1)", "1.0 if known exploited in the wild; otherwise max(EPSS probability, 0.5 if public exploit exists)"),
        ("Asset (0-1)", "0.5 × data classification (Restricted 1.0, Confidential 0.7, Internal 0.4, Public 0.1) + 0.3 × internet facing + 0.2 × environment (Prod 1.0, Test 0.4, Dev 0.3)"),
        ("Priority", "P1 ≥ 70, P2 ≥ 55, P3 ≥ 35, else P4. Known-exploited on an internet-facing asset is always P1."),
        ("SLA", ", ".join(f"{k}: {v} days" for k, v in SLA_DAYS.items()) + " from first detection."),
        ("SLA compliance", "Share of remediated findings closed within their SLA window."),
        ("Overdue", "Open findings whose SLA due date has passed as of the report date."),
        ("Data", "All hosts, owners and vulnerability IDs (SYN-*) are synthetic. No real company data is used."),
    ]
    _write_table(ws, pd.DataFrame(lines, columns=["Item", "Definition"]), start_row=4)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120


def write_workbook(r: Report, path: Path) -> Path:
    wb = Workbook()
    summary_sheet(wb, r)
    findings_sheet(wb, "Open Findings (Prioritized)", r.open)
    findings_sheet(wb, "Overdue", r.overdue)
    simple_sheet(
        wb, "Severity vs Priority", "CVSS Severity vs Risk-Based Priority", r.reprioritization.reset_index(),
        "Rows: what the scanner says. Columns: what the risk model says after adding threat intel and asset context.",
        chart={"title": "Re-prioritization of Open Findings", "stacked": True},
    )
    simple_sheet(wb, "By Business Unit", "Open Risk by Business Unit", r.by_business_unit,
                 "Sorted by open P1, then P1+P2. SLA compliance uses findings remediated in the period.",
                 chart={"title": "Open P1/P2 by Business Unit", "max_col": 3, "stacked": True, "anchor": "M4"})
    simple_sheet(wb, "Top Risky Assets", "Top 15 Assets by Total Open Risk", r.top_assets, "Sorted by P1 count, then highest single risk. Start remediation conversations with these owners.")
    simple_sheet(wb, "Quick Wins", "Top 10 Remediation Quick Wins", r.quick_wins,
                 "Each row is one patch or configuration change, ranked by the P1/P2 risk it removes across all affected assets.")
    simple_sheet(wb, "Aging", "Open Finding Age by Priority", r.aging, "Old P1/P2 findings signal a process gap.",
                 chart={"title": "Aging of Open Findings", "stacked": True})
    simple_sheet(wb, "MTTR", "Mean Time to Remediate (Fixed Findings)", r.mttr, "Days from first detection to fix.")
    methodology_sheet(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
