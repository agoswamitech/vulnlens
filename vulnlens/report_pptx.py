"""Build a short executive PowerPoint deck from the analysis."""

from __future__ import annotations

import tempfile
from pathlib import Path

import matplotlib
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from .analysis import PRIORITY_ORDER, SEVERITY_ORDER, Report  # noqa: E402
from .scoring import SLA_DAYS  # noqa: E402

NAVY = RGBColor(0x1F, 0x38, 0x64)
INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x52, 0x51, 0x4E)
TILE = RGBColor(0xF2, 0xF5, 0xFA)

# Priority is ordinal, so it uses one sequential hue (darkest = most urgent)
PRIORITY_COLORS = {"P1": "#104281", "P2": "#2a78d6", "P3": "#6da7ec", "P4": "#b7d3f6"}
SERIES_1, SERIES_2 = "#2a78d6", "#eb6834"
GRID, AXIS_TEXT = "#e4e3df", "#52514e"


def _style(ax) -> None:
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#bdbcb6")
    ax.tick_params(colors=AXIS_TEXT, length=0, labelsize=10)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def _save(fig, tmp: Path, name: str) -> Path:
    path = tmp / f"{name}.png"
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def chart_reprioritization(r: Report, tmp: Path) -> Path:
    ct = r.reprioritization.loc[[s for s in SEVERITY_ORDER if r.reprioritization.loc[s].sum() > 0]]
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    bottom = [0] * len(ct)
    for p in PRIORITY_ORDER:
        vals = ct[p].tolist()
        ax.bar(ct.index, vals, bottom=bottom, color=PRIORITY_COLORS[p], label=p, width=0.55, edgecolor="white", linewidth=1.5)
        bottom = [b + v for b, v in zip(bottom, vals)]
    for i, total in enumerate(bottom):
        ax.text(i, total + max(bottom) * 0.02, f"{total}", ha="center", color="#0b0b0b", fontsize=10)
    _style(ax)
    ax.set_xlabel("Scanner (CVSS) severity", color=AXIS_TEXT)
    ax.set_ylabel("Open findings", color=AXIS_TEXT)
    ax.legend(title="Risk priority", frameon=False, fontsize=9, title_fontsize=9, loc="upper right")
    return _save(fig, tmp, "reprioritization")


def chart_trend(r: Report, tmp: Path) -> Path:
    t = r.trend
    labels = [pd.Period(m).strftime("%b") for m in t["Month"]]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(9.6, 3.4))
    x = range(len(t))
    a1.bar([i - 0.19 for i in x], t["Opened"], width=0.36, color=SERIES_1, label="Opened")
    a1.bar([i + 0.19 for i in x], t["Closed"], width=0.36, color=SERIES_2, label="Closed")
    a1.set_xticks(list(x), labels)
    a1.set_title("Findings opened vs closed per month", loc="left", fontsize=11, color="#0b0b0b", pad=22)
    a1.legend(frameon=False, fontsize=9, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2)
    a1.set_ylim(0, max(t["Opened"].max(), t["Closed"].max()) * 1.15)
    _style(a1)
    a2.plot(list(x), t["Backlog"], color=SERIES_1, linewidth=2, marker="o", markersize=6)
    a2.annotate(f"{t['Backlog'].iloc[-1]}", (len(t) - 1, t["Backlog"].iloc[-1]), textcoords="offset points", xytext=(0, 8), ha="center", fontsize=10)
    a2.set_xticks(list(x), labels)
    a2.set_ylim(0, t["Backlog"].max() * 1.2)
    a2.set_title("Unresolved backlog at month end", loc="left", fontsize=11, color="#0b0b0b", pad=22)
    _style(a2)
    fig.tight_layout(w_pad=3)
    return _save(fig, tmp, "trend")


def chart_business_units(r: Report, tmp: Path) -> Path:
    bu = r.by_business_unit.iloc[::-1]
    fig, ax = plt.subplots(figsize=(5.6, 3.6))
    left = [0] * len(bu)
    for p in ("P1", "P2"):
        ax.barh(bu["Business Unit"], bu[p], left=left, color=PRIORITY_COLORS[p], label=p, height=0.55, edgecolor="white", linewidth=1.5)
        left = [a + b for a, b in zip(left, bu[p])]
    _style(ax)
    ax.yaxis.grid(False)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_xlabel("Open P1 + P2 findings", color=AXIS_TEXT)
    ax.legend(frameon=False, fontsize=9, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2)
    return _save(fig, tmp, "business_units")


def chart_quick_wins(r: Report, tmp: Path) -> Path:
    q = r.quick_wins.head(5).iloc[::-1]
    fig, ax = plt.subplots(figsize=(6.2, 3.4))
    ax.barh(q["title"].str.slice(0, 42), q["Share_of_P1_P2_Risk_%"], color=SERIES_1, height=0.55)
    for i, (v, n) in enumerate(zip(q["Share_of_P1_P2_Risk_%"], q["Assets"])):
        ax.text(v + 0.2, i, f"{v:.1f}%  ({n} assets)", va="center", fontsize=9, color="#0b0b0b")
    _style(ax)
    ax.yaxis.grid(False)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_xlim(0, q["Share_of_P1_P2_Risk_%"].max() * 1.45)
    ax.set_xlabel("Share of open P1/P2 risk removed", color=AXIS_TEXT)
    return _save(fig, tmp, "quick_wins")


# ---------- slide helpers ----------

def _text(slide, left, top, width, height, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.alignment = line, align
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        p.space_after = Pt(6)
    return box


def _slide(prs, title: str, subtitle: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    _text(s, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7), title, size=26, bold=True, color=NAVY)
    if subtitle:
        _text(s, Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.5), subtitle, size=14, color=MUTED)
    return s


def _tile(slide, left, top, value: str, label: str):
    shp = slide.shapes.add_shape(1, left, top, Inches(2.9), Inches(1.45))
    shp.fill.solid()
    shp.fill.fore_color.rgb = TILE
    shp.line.fill.background()
    _text(slide, left + Inches(0.15), top + Inches(0.1), Inches(2.6), Inches(0.7), value, size=32, bold=True, color=NAVY)
    _text(slide, left + Inches(0.15), top + Inches(0.8), Inches(2.6), Inches(0.6), label, size=12, color=MUTED)


def takeaways(r: Report) -> list[str]:
    k = r.kpis
    top_bu = r.by_business_unit.iloc[0]
    qw = r.quick_wins.iloc[0]
    return [
        f"{k['Open P1 (fix within 15 days)']} findings need action within 15 days; {k['Open known-exploited findings']} are flaws attackers are known to exploit.",
        f"Risk-based scoring moved {k['CVSS Critical/High downgraded to P3/P4']} scanner Critical/High findings down and "
        f"{k['CVSS Medium/Low escalated to P1/P2']} Medium/Low findings up, focusing effort where the data says it matters.",
        f"{k['Overdue (past SLA)']} open findings are past their SLA; {k['SLA compliance (fixed on time)']:.0f}% of fixes this period closed on time.",
        f"{top_bu['Business Unit']} has the most urgent work: {int(top_bu['P1'])} P1 and {int(top_bu['P2'])} P2 findings across {int(top_bu['Assets'])} assets.",
        f"Five fixes would remove {k['P1/P2 risk removed by top 5 fixes (%)']:.0f}% of urgent (P1/P2) risk; the largest is \"{qw['title']}\" on {int(qw['Assets'])} assets.",
    ]


def write_deck(r: Report, path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)

        # 1. Title
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY
        bg.line.fill.background()
        white = RGBColor(0xFF, 0xFF, 0xFF)
        _text(s, Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.2), "Vulnerability Risk Review", size=44, bold=True, color=white)
        _text(s, Inches(0.8), Inches(3.7), Inches(11.5), Inches(0.8), f"Risk-based prioritization and remediation performance · {r.as_of:%B %Y}", size=20, color=white)
        _text(s, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.5), "Synthetic data for demonstration. No real company systems are represented.", size=12, color=RGBColor(0xC8, 0xD3, 0xE6))

        # 2. Executive summary
        s = _slide(prs, "Executive summary", f"{r.kpis['Open findings']} open findings across {r.kpis['Assets with open findings']} assets")
        k = r.kpis
        tiles = [
            (str(k["Open P1 (fix within 15 days)"]), "Open P1 findings (15-day SLA)"),
            (str(k["Open known-exploited findings"]), "Known-exploited flaws open"),
            (str(k["Overdue (past SLA)"]), "Open findings past SLA"),
            (f"{k['SLA compliance (fixed on time)']:.0f}%", "Fixes closed within SLA"),
        ]
        for i, (v, lbl) in enumerate(tiles):
            _tile(s, Inches(0.5 + i * 3.1), Inches(1.6), v, lbl)
        _text(s, Inches(0.5), Inches(3.4), Inches(12.3), Inches(3.8), ["• " + t for t in takeaways(r)], size=15)

        # 3. Severity is not risk
        s = _slide(prs, "Scanner severity is not the same as risk", "Adding threat intelligence and asset context changes what gets fixed first")
        s.shapes.add_picture(str(chart_reprioritization(r, tmp)), Inches(0.5), Inches(1.6), height=Inches(5.2))
        _text(s, Inches(9.2), Inches(1.8), Inches(3.7), Inches(5), [
            "Risk score blends three inputs:",
            "• Severity (CVSS), 40%",
            "• Threat: known exploitation or EPSS likelihood, 30%",
            "• Asset: data sensitivity, internet exposure, production, 30%",
            f"{k['CVSS Medium/Low escalated to P1/P2']} Medium/Low findings were escalated, typically exploitable flaws on systems holding policyholder PII.",
        ], size=13)

        # 4. Remediation performance
        s = _slide(prs, "Remediation performance", f"SLA targets: " + ", ".join(f"{p} {d} days" for p, d in SLA_DAYS.items()))
        s.shapes.add_picture(str(chart_trend(r, tmp)), Inches(0.5), Inches(1.5), width=Inches(12.3))
        mttr = r.mttr.dropna()
        _text(s, Inches(0.5), Inches(5.9), Inches(12.3), Inches(1.2),
              "Median days to fix:  " + "   ·   ".join(f"{row.Severity} {row.MTTR_Median_Days:.0f}" for row in mttr.itertuples()), size=14, color=MUTED)

        # 5. Where risk lives
        s = _slide(prs, "Where the risk lives", "Open P1 and P2 findings by business unit, and the assets to start with")
        s.shapes.add_picture(str(chart_business_units(r, tmp)), Inches(0.5), Inches(1.6), height=Inches(4.9))
        top = r.top_assets.head(6)
        _text(s, Inches(8.2), Inches(1.6), Inches(4.8), Inches(0.5), "Highest-risk assets", size=15, bold=True, color=NAVY)
        _text(s, Inches(8.2), Inches(2.1), Inches(4.8), Inches(4.8), [
            f"{a.hostname.split('.')[0]} · {a.business_unit} · {a.Open} open, {a.P1} P1" + (" · internet-facing" if a.internet_facing else "")
            for a in top.itertuples()
        ], size=12)

        # 6. Quick wins
        s = _slide(prs, "Quick wins: five fixes, outsized impact", "Each bar is one patch or configuration change; share of all open P1/P2 risk it removes")
        s.shapes.add_picture(str(chart_quick_wins(r, tmp)), Inches(0.5), Inches(1.6), height=Inches(5.0))

        # 7. Recommendations
        s = _slide(prs, "Recommendations")
        _text(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.6), [
            "1. Patch all known-exploited findings on internet-facing systems this week, starting with the quick-wins list.",
            "2. Adopt risk-based SLAs (P1-P4) in place of CVSS-only SLAs so teams spend effort where exploitation is likely.",
            f"3. Review overdue items with the {r.by_business_unit.iloc[0]['Business Unit']} and {r.by_business_unit.iloc[1]['Business Unit']} owners in the next risk meeting.",
            f"4. Re-validate the {k['Risk accepted (exceptions)']} risk acceptances quarterly, with an expiry date and compensating control for each.",
            "5. Track SLA compliance and backlog monthly as program KPIs and report them to IT leadership.",
        ], size=17)

        # 8. Appendix
        s = _slide(prs, "Appendix: methodology")
        _text(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.6), [
            "Risk score = 100 × (0.40 × CVSS/10 + 0.30 × Threat + 0.30 × Asset)",
            "Threat = 1.0 if known exploited, else max(EPSS, 0.5 if a public exploit exists)",
            "Asset = 0.5 × data classification + 0.3 × internet facing + 0.2 × environment",
            "Priority: P1 ≥ 70, P2 ≥ 55, P3 ≥ 35, else P4; known-exploited and internet-facing is always P1",
            "SLA compliance = remediated findings closed within their SLA window",
            "Full finding-level detail is in the accompanying Excel workbook.",
        ], size=15)

        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)
    return path
